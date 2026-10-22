import io
import json
import socket
from pathlib import Path

import pytest
from docx import Document
from fastapi.testclient import TestClient

from knowledge import KnowledgeBase, parse_file
from runtime import OfflineRuntime
from stateflow_gate import parse_and_validate_state


@pytest.fixture
def kb(tmp_path, monkeypatch):
    def deny_network(*args, **kwargs):
        raise AssertionError('Offline tests must not connect to network')
    monkeypatch.setattr(socket.socket, 'connect', deny_network)
    db = KnowledgeBase(tmp_path / 'db')
    yield db
    db.close()


def test_ingest_search_gate_and_persist(kb):
    kb.ingest('制度.md', '差旅报销需要有效发票、行程单及费用明细。员工应在十五日内提交。'.encode())
    result = OfflineRuntime(kb).ask('差旅报销材料')
    assert result['gate']['ok']
    assert result['citations']
    assert all(c['text'] in '差旅报销需要有效发票、行程单及费用明细。员工应在十五日内提交。' for c in result['citations'])
    assert result['evidence_state']['decision'] == 'search'
    assert result['mode'] == 'offline_extract_only'
    assert not kb.search('火星飞船推进器')
    path = kb.directory
    kb.qdrant.close()
    kb.sql.close()
    restored = KnowledgeBase(path)
    try:
        assert restored.search('差旅报销')
    finally:
        restored.close()
    # fixture close is idempotent for Qdrant/sqlite


def test_replace_and_deduplicate(kb):
    first = kb.ingest('制度.txt', '旧版制度住宿标准为五百元。'.encode())
    assert kb.ingest('制度.txt', '旧版制度住宿标准为五百元。'.encode())['unchanged']
    new = kb.ingest('制度.txt', '新版制度住宿标准为六百元。'.encode())
    assert first['version'] != new['version']
    assert len(kb.documents()) == 1
    assert all('旧版' not in h['text'] for h in kb.search('住宿标准'))


def test_cross_document_ids_and_chunking(kb):
    text = '\n'.join(f'第{i}条，费用报销必须提供发票与审批材料。' for i in range(100))
    assert kb.ingest('a.md', text.encode())['chunks'] > 1
    kb.ingest('b.md', text.encode())
    ids = [s['span_id'] for r in kb.rows for s in json.loads(r['spans'])]
    assert len(ids) == len(set(ids))


def test_gate_rejects_fabricated_evidence():
    state = {'evidence': [{'span_id': 'd:s1', 'text': '编造的报销规定', 'role': 'direct'}],
             'open_questions': [], 'decision': 'answer'}
    result = parse_and_validate_state('<evidence_state>' + json.dumps(state) + '</evidence_state>',
                                     {'d:s1': {'text': '实际原文'}})
    assert not result.ok


def test_parsers_and_invalid_files():
    document = Document()
    document.add_paragraph('报销材料包含发票。')
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = '项目'
    table.cell(0, 1).text = '限额'
    buffer = io.BytesIO()
    document.save(buffer)
    parsed = parse_file('制度.docx', buffer.getvalue())
    assert any('报销材料' in t for _, t in parsed)
    assert any('项目 | 限额' in t for _, t in parsed)
    with pytest.raises(ValueError):
        parse_file('bad.exe', b'123')
    with pytest.raises(ValueError):
        parse_file('empty.txt', b'')


def test_http_end_to_end(tmp_path, monkeypatch):
    monkeypatch.setenv('KNOWLEDGE_DATA_DIR', str(tmp_path / 'api'))
    from app import app
    with TestClient(app) as client:
        assert client.get('/health').json()['llm_configured'] is False
        assert client.get('/').status_code == 200
        assert client.post('/demo').status_code == 200
        assert client.post('/ingest', files={'file': ('upload.txt', '发票需要真实有效。'.encode())}).status_code == 200
        response = client.post('/retrieve', json={'queries': ['差旅报销'], 'topk': 3, 'return_scores': True})
        assert response.status_code == 200
        doc = response.json()['result'][0][0]['document']
        assert doc['spans'] and doc['doc_id']
        answer = client.post('/ask', json={'question': '差旅报销需要什么材料？'}).json()
        assert answer['gate']['ok'] and answer['citations']
        assert client.post('/retrieve', json={'queries': [''], 'topk': 3}).status_code == 400
        assert client.post('/retrieve', json={'queries': ['a'], 'topk': 999}).status_code == 422
